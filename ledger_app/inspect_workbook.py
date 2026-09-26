#!/usr/bin/env python3
"""Structure report for Master Ledger.xlsm — every sheet, every formula, every macro.

This is how the ledger's layout and VBA can be read for porting the macros to
the web app, without the workbook itself (it holds real customer and payment
data) going anywhere. It only reads the file, and nothing is committed to this
repo, which is public and serves the live apps:

    bash ledger_app/report.sh                # find the ledger, write ledger_app/out/
    bash ledger_app/report.sh --push         # ...and send it to your private repo
    python3 -m ledger_app.inspect_workbook "/path/to/Master Ledger.xlsm"

For each workbook it writes ``ledger_app/out/<name>.md`` (to read) and
``<name>.json`` (the same, for code):

* **every sheet** — used range, code name, hidden/protected, freeze panes, Excel
  tables (with calculated columns), merged cells, drop-down lists, conditional
  formats, the input cells left unlocked on protected sheets, print area and
  titles, hidden rows/columns, comments, charts/images/pivots, and the buttons
  on it with the macro each one runs;
* the cells themselves — the whole sheet for small sheets (forms, dashboards),
  the top rows for big ones — plus the header rows found anywhere on the sheet,
  a profile of every column and a few sample rows;
* **every distinct formula** on the sheet (a formula filled down a column counts
  once, with the range it covers);
* defined names and links to other workbooks;
* **the full VBA source** of every module, with each macro's entry point,
  buttons, dialogs and InputBox prompts, the event macros Excel runs by itself,
  and the helper procedures they call.

Sheets that share a layout (one ledger per customer) are described in full once
and then listed as "same layout as …", with any formulas that differ.

Values are **masked by default**: numbers become ``#``; phone, GSTIN, PAN,
account numbers and e-mail addresses in text are blurred; string literals on
password-looking VBA lines become ``"***"``. Formulas, headers, labels, number
formats and dates are kept — that's the structure. ``--full-values`` turns the
value masking off.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import plistlib
import posixpath
import re
import shutil
import subprocess
import sys
import zipfile
from urllib.parse import unquote

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from ledger_app import grid as G   # noqa: E402
from ledger_app import vba as V    # noqa: E402

TOOL_VERSION = 1
OUT_DIR = os.path.join(_HERE, "out")   # git-ignored (ledger_app/.gitignore)
MAX_ROWS = 20000       # rows read per sheet
MAX_COLS = 250         # columns read per sheet
TOP_ROWS = 12          # cells listed for big sheets
SMALL_SHEET = (100, 40)  # sheets up to this many rows x cols are listed cell by cell
PROFILE_COLS = 80
MAX_PATTERNS = 250     # distinct formulas listed per sheet

VBA_EXTS = (".xlsm", ".xlsb", ".xls", ".xlam", ".xla", ".xltm")
LEDGER_NAME = "Master Ledger.xlsm"
AGENT_PLIST = os.path.expanduser("~/Library/LaunchAgents/com.vriddhi.paymentagent.plist")
ICLOUD = os.path.expanduser("~/Library/Mobile Documents/com~apple~CloudDocs")


# ---- finding the ledger on the Mac ------------------------------------------------

def find_ledger() -> str | None:
    """Where Master Ledger.xlsm is: $MASTER_LEDGER_PATH, else the payment agent's
    configured path, else the usual iCloud folders, else a Spotlight search."""
    cands = [os.environ.get("MASTER_LEDGER_PATH", "")]
    try:
        with open(AGENT_PLIST, "rb") as fh:
            cands.append(plistlib.load(fh).get("EnvironmentVariables", {}).get("MASTER_LEDGER_PATH", ""))
    except (OSError, ValueError, plistlib.InvalidFileException):
        pass
    cands += [os.path.join(ICLOUD, "Vriddhi Fuels", LEDGER_NAME), os.path.join(ICLOUD, LEDGER_NAME)]
    for c in cands:
        c = os.path.expanduser(c or "")
        if c and os.path.exists(c):
            return c
    if shutil.which("mdfind"):
        try:
            out = subprocess.run(["mdfind", f'kMDItemFSName == "{LEDGER_NAME}"'],
                                 capture_output=True, text=True, timeout=30).stdout
            hits = [ln for ln in out.splitlines() if ln.strip() and "/.Trash/" not in ln]
            if hits:
                return hits[0]
        except (OSError, subprocess.SubprocessError):
            pass
    return None


def _materialize(path: str) -> None:
    """Make sure an iCloud file is downloaded, not just a placeholder."""
    if sys.platform == "darwin" and shutil.which("brctl"):
        try:
            subprocess.run(["brctl", "download", path], capture_output=True, timeout=120)
        except (OSError, subprocess.SubprocessError):
            pass


# ---- package parts: buttons and external links --------------------------------

def _unescape(s: str) -> str:
    return (s or "").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"') \
        .replace("&apos;", "'").replace("&amp;", "&")


def _rels(z: zipfile.ZipFile, part: str) -> dict:
    """Relationship id -> (type, absolute target) for a package part."""
    d, b = posixpath.split(part)
    path = posixpath.join(d, "_rels", b + ".rels")
    if path not in z.namelist():
        return {}
    xml = z.read(path).decode("utf-8", "replace")
    out = {}
    for m in re.finditer(r"<Relationship\b([^>]*)/?>", xml):
        attrs = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        if attrs.get("TargetMode") == "External":
            continue
        tgt = attrs.get("Target", "")
        full = posixpath.normpath(tgt.lstrip("/") if tgt.startswith("/") else posixpath.join(d, tgt))
        out[attrs.get("Id")] = (attrs.get("Type", "").rsplit("/", 1)[-1], full)
    return out


def find_buttons(path: str) -> list:
    """[{sheet, text, macro, kind}] for every control or shape that runs a macro."""
    out = []
    try:
        z = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile):
        return out
    with z:
        names = set(z.namelist())
        if "xl/workbook.xml" not in names:
            return out
        wb_xml = z.read("xl/workbook.xml").decode("utf-8", "replace")
        wb_rels = _rels(z, "xl/workbook.xml")
        for m in re.finditer(r"<sheet\b([^>]*)/?>", wb_xml):
            attrs = dict(re.findall(r'([\w:]+)="([^"]*)"', m.group(1)))
            rid = attrs.get("r:id")
            if rid not in wb_rels:
                continue
            sheet = _unescape(attrs.get("name", ""))
            for kind, target in _rels(z, wb_rels[rid][1]).values():
                if target not in names:
                    continue
                xml = z.read(target).decode("utf-8", "replace")
                if kind == "vmlDrawing":    # form controls (buttons, check boxes …)
                    for shp in re.finditer(r"<v:shape\b.*?</v:shape>", xml, re.S):
                        block = shp.group(0)
                        mac = re.search(r"<x:FmlaMacro>(.*?)</x:FmlaMacro>", block, re.S)
                        if not mac:
                            continue
                        typ = re.search(r'<x:ClientData\s+ObjectType="(\w+)"', block)
                        txt = re.search(r"<v:textbox\b.*?>(.*?)</v:textbox>", block, re.S)
                        text = " ".join(_unescape(re.sub(r"<[^>]+>", " ", txt.group(1))).split()) if txt else ""
                        out.append({"sheet": sheet, "macro": _unescape(mac.group(1).strip()),
                                    "text": text, "kind": typ.group(1) if typ else "control"})
                elif kind == "drawing":     # shapes / pictures with "Assign Macro"
                    for shp in re.finditer(r"<xdr:(sp|pic|grpSp|graphicFrame)\b([^>]*)>(.*?)</xdr:\1>",
                                           xml, re.S):
                        mac = re.search(r'\bmacro="([^"]+)"', shp.group(2))
                        if not mac:
                            continue
                        name = re.search(r'<xdr:cNvPr\b[^>]*\bname="([^"]*)"', shp.group(3))
                        texts = re.findall(r"<a:t>(.*?)</a:t>", shp.group(3), re.S)
                        text = _unescape(" ".join(texts).strip()) or (_unescape(name.group(1)) if name else "")
                        out.append({"sheet": sheet, "macro": _unescape(mac.group(1)),
                                    "text": text, "kind": "shape"})
    return out


def external_links(path: str) -> list:
    """Other workbooks this one's formulas point at."""
    out = []
    try:
        with zipfile.ZipFile(path) as z:
            for n in sorted(z.namelist()):
                if re.match(r"xl/externalLinks/_rels/externalLink\d+\.xml\.rels$", n):
                    xml = z.read(n).decode("utf-8", "replace")
                    out += [unquote(_unescape(t)) for t in re.findall(r'Target="([^"]*)"', xml)]
    except (OSError, zipfile.BadZipFile):
        pass
    return out


# ---- masking / display --------------------------------------------------------------

_EMAIL = re.compile(r"\b([A-Za-z0-9])[A-Za-z0-9._%+-]*@([A-Za-z0-9.-]+\.[A-Za-z]{2,})\b")
_PAN = re.compile(r"\b([A-Z]{5})(\d{4})([A-Z])\b")


def mask_text(s: str) -> str:
    """Blur phone / GSTIN / PAN / account numbers and e-mail addresses."""
    s = _EMAIL.sub(r"\1***@\2", s)
    s = _PAN.sub(lambda m: m.group(1) + "9999" + m.group(3), s)
    if sum(ch.isdigit() for ch in s) >= 7:
        s = re.sub(r"\d", "9", s)
    return s


def show(value, formula, mask: bool) -> str:
    """How a cell appears in the report."""
    if formula:
        return formula if len(formula) <= 160 else formula[:159] + "…"
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.strftime("%Y-%m-%d")
    if isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    if isinstance(value, bool):
        return str(value).upper()
    if G.is_num(value):
        return "#" if mask else str(G.display(value))
    s = " ⏎ ".join(str(value).split("\n")).strip()
    if mask:
        s = mask_text(s)
    return s if len(s) <= 60 else s[:59] + "…"


def _kind(value, formula) -> str:
    if formula:
        return "formula"
    if G.is_blank(value):
        return ""
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return "date"
    if isinstance(value, bool):
        return "bool"
    if G.is_num(value):
        return "number"
    return "text"


# ---- formulas -------------------------------------------------------------------

_REF = re.compile(r"(?<![A-Za-z_\d.\]$])(\$?)([A-Z]{1,3})(\$?)(\d{1,7})(?![A-Za-z_\d(\[!])")
_QUOTED = re.compile(r'("(?:[^"]|"")*"|\'(?:[^\']|\'\')*\')')


def to_r1c1(formula: str, row: int, col: int) -> str:
    """A1 references -> R1C1 relative to the cell at (row, col), so the same
    formula filled down a column reads the same in every row."""
    def rc(m):
        cabs, letters, rabs, digits = m.groups()
        c, r = G.col_index(letters), int(digits) - 1
        rs = f"R{r + 1}" if rabs else ("R" if r == row else f"R[{r - row}]")
        cs = f"C{c + 1}" if cabs else ("C" if c == col else f"C[{c - col}]")
        return rs + cs
    parts = _QUOTED.split(formula)
    return "".join(p if i % 2 else _REF.sub(rc, p) for i, p in enumerate(parts))


def _span(first, last, count) -> str:
    if count == 1:
        return G.a1(*first)
    if first[1] == last[1] or first[0] == last[0]:
        return f"{G.a1(*first)}:{G.a1(*last)}"
    return f"{G.a1(*first)} … {G.a1(*last)}"


def formula_patterns(grid: G.Grid) -> tuple:
    """(distinct formulas in order of first appearance, total formula cells)."""
    pats, order, total = {}, [], 0
    for r, row in enumerate(grid.formulas or []):
        for c, f in enumerate(row or []):
            if not f:
                continue
            total += 1
            key = to_r1c1(f, r, c)
            p = pats.get(key)
            if p is None:
                p = pats[key] = {"key": key, "formula": f, "first": (r, c), "last": (r, c), "count": 0}
                order.append(p)
            p["count"] += 1
            p["last"] = (r, c)
    return ([{"key": p["key"], "formula": p["formula"], "count": p["count"],
              "cells": _span(p["first"], p["last"], p["count"])} for p in order], total)


# ---- sheet analysis -----------------------------------------------------------------

def _headerish(v) -> bool:
    """Could be a column header: short text, not a date, not an ID or account
    number (labels like "GSTIN: 21AAB…" or "A/c 5921…" are data)."""
    if not G.is_label(v):
        return False
    s = str(v).strip()
    return len(s) <= 32 and sum(ch.isdigit() for ch in s) < 5 and G.to_date(s) is None


def _starts_section(grid: G.Grid, r: int, labels: list) -> bool:
    """A header row starts a region: the row above it is empty over the same
    columns, holds a single title, or is a group-header row with fewer labels
    (e.g. "Stock" merged over Opening | Receipt | Sales). A data row of a log
    fails this — the row above it is the header or another data row."""
    if r == 0:
        return True
    lo, hi = labels[0][0], labels[-1][0]
    filled = [(grid.v(r - 1, c), grid.f(r - 1, c)) for c in range(lo, hi + 1)
              if grid.f(r - 1, c) or not G.is_blank(grid.v(r - 1, c))]
    if len(filled) <= 1:
        return True
    return all(not f and _headerish(v) for v, f in filled) and len(filled) < len(labels)


def _range_bounds(ref: str):
    """"A1:M32" -> 0-based (r1, c1, r2, c2), or None."""
    m = re.fullmatch(r"\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?", (ref or "").strip())
    if not m:
        return None
    c1, r1 = G.col_index(m.group(1)), int(m.group(2)) - 1
    c2, r2 = (G.col_index(m.group(3)), int(m.group(4)) - 1) if m.group(3) else (c1, r1)
    return r1, c1, r2, c2


def header_blocks(grid: G.Grid, scan_rows: int = 60, cap: int = 14, mask: bool = True,
                  tables: list | None = None) -> list:
    """Rows that look like column headers: 3+ short labels side by side (one
    empty cell allowed between, for merged headers) that start a region. Past
    the first ``scan_rows`` rows only rows under a blank row are checked, and
    rows inside an Excel table's data are never headers (the table's own header
    row is). Labels are masked like any other cell text."""
    bodies = []
    for t in tables or []:
        b = _range_bounds(t.get("ref", ""))
        if b:
            bodies.append((b[0] + 1, b[1], b[2], b[3]))
    blocks, index = [], {}
    for r in range(grid.nrows):
        if r >= scan_rows and any(not G.is_blank(x) for x in (grid.values[r - 1] or [])):
            continue
        c = 0
        while c < grid.ncols:
            if not _headerish(grid.v(r, c)):
                c += 1
                continue
            labels = []
            while c < grid.ncols:
                if _headerish(grid.v(r, c)):
                    labels.append((c, str(grid.v(r, c)).strip()))
                    c += 1
                elif G.is_blank(grid.v(r, c)) and _headerish(grid.v(r, c + 1)):
                    c += 1
                else:
                    break
            if len(labels) < 3 or not _starts_section(grid, r, labels):
                continue
            lo, hi = labels[0][0], labels[-1][0]
            if any(r1 <= r <= r2 and lo <= c2 and hi >= c1 for r1, c1, r2, c2 in bodies):
                continue
            key = tuple((cc, G.norm(t)) for cc, t in labels)
            if key in index:
                blocks[index[key]]["repeats"].append(r + 1)
                continue
            index[key] = len(blocks)
            blocks.append({"row": r + 1, "range": f"{G.a1(r, lo)}:{G.a1(r, hi)}",
                           "first_col": lo, "last_col": hi,
                           "labels": [{"col": G.col_letter(cc), "label": mask_text(t) if mask else t}
                                      for cc, t in labels],
                           "repeats": [], "_key": key})
            if len(blocks) >= cap:
                return blocks
    return blocks


def column_profiles(grid: G.Grid, fmts: list, blocks: list) -> list:
    """Per column: header, how many cells of each type, number format, formulas."""
    out = []
    for c in range(min(grid.ncols, PROFILE_COLS)):
        hdr_row, header = -1, ""
        for b in blocks:
            hit = [x["label"] for x in b["labels"] if x["col"] == G.col_letter(c)]
            if hit:
                hdr_row, header = b["row"] - 1, hit[0]
                break
        counts = {"number": 0, "date": 0, "text": 0, "formula": 0, "bool": 0}
        first_f = last_f = fmt = ""
        last_row = 0
        for r in range(hdr_row + 1, grid.nrows):
            v, f = grid.v(r, c), grid.f(r, c)
            k = _kind(v, f)
            if not k:
                continue
            counts[k] += 1
            last_row = r + 1
            if f:
                first_f = first_f or f
                last_f = f
            if not fmt and r < len(fmts) and c < len(fmts[r]) and fmts[r][c] not in (None, "General"):
                fmt = fmts[r][c]
        filled = sum(counts.values())
        if not filled and not header:
            continue
        out.append({"col": G.col_letter(c), "header": header,
                    "header_row": hdr_row + 1 if hdr_row >= 0 else None,
                    "filled": filled, "types": {k: v for k, v in counts.items() if v},
                    "number_format": fmt, "first_formula": first_f,
                    "last_formula": last_f if last_f != first_f else "", "last_row": last_row})
    return out


def cell_rows(grid: G.Grid, mask: bool, upto: int) -> list:
    """Non-empty cells, row by row: [{"row": 12, "cells": [["B", "Bill To"], …]}]."""
    out = []
    for r in range(min(grid.nrows, upto)):
        cells = [[G.col_letter(c), s] for c in range(grid.ncols)
                 for s in [show(grid.v(r, c), grid.f(r, c), mask)] if s != ""]
        if cells:
            out.append({"row": r + 1, "cells": cells})
    return out


def _samples(grid: G.Grid, blocks: list, mask: bool) -> list:
    """First three and last two rows under the first header row."""
    if not blocks:
        return []
    b = blocks[0]
    lo, hi = b["first_col"], min(b["last_col"], b["first_col"] + 24)
    rows = [r for r in range(b["row"], grid.nrows)
            if any(not G.is_blank(grid.v(r, c)) or grid.f(r, c) for c in range(lo, hi + 1))]
    pick = rows[:3] + [r for r in rows[-2:] if r not in rows[:3]]
    return [{"row": r + 1, "cells": [[G.col_letter(c), s] for c in range(lo, hi + 1)
                                     for s in [show(grid.v(r, c), grid.f(r, c), mask)] if s]}
            for r in pick]


def _title(grid: G.Grid, mask: bool) -> str:
    texts = [show(grid.v(r, c), None, mask) for r in range(min(grid.nrows, 4))
             for c in range(min(grid.ncols, 12)) if isinstance(grid.v(r, c), str) and grid.v(r, c).strip()]
    return " · ".join(texts[:2])


def _ranges(cells: list, cap: int = 60) -> tuple:
    """0-based (row, col) cells -> compact A1 ranges (runs across, then down)."""
    runs = {}
    by_row = {}
    for r, c in cells:
        by_row.setdefault(r, set()).add(c)
    for r in sorted(by_row):
        cs = sorted(by_row[r])
        start = prev = cs[0]
        for c in cs[1:] + [None]:
            if c is not None and c == prev + 1:
                prev = c
                continue
            runs.setdefault((start, prev), []).append(r)
            if c is not None:
                start = prev = c
    out = []
    for (c1, c2), rows in runs.items():
        top = last = rows[0]
        for r in rows[1:] + [None]:
            if r is not None and r == last + 1:
                last = r
                continue
            a, b = G.a1(top, c1), G.a1(last, c2)
            out.append((top, c1, a if a == b else f"{a}:{b}"))
            if r is not None:
                top = last = r
    out.sort()
    return [x[2] for x in out[:cap]], len(out)


def _hidden(ws) -> dict:
    cols, rows = [], []
    for key, d in ws.column_dimensions.items():
        if d.hidden:
            lo = d.min or (G.col_index(key) + 1)
            hi = d.max or lo
            cols.append(G.col_letter(lo - 1) if lo == hi else f"{G.col_letter(lo - 1)}:{G.col_letter(hi - 1)}")
    for r, d in sorted(ws.row_dimensions.items()):
        if d.hidden:
            rows.append(r)
    spans, start, prev = [], None, None
    for r in rows + [None]:
        if r is not None and prev is not None and r == prev + 1:
            prev = r
            continue
        if start is not None:
            spans.append(str(start) if start == prev else f"{start}-{prev}")
        start = prev = r
    return {"columns": cols, "rows": spans}


def _widths(ws) -> dict:
    out = {}
    for key, d in ws.column_dimensions.items():
        if d.width and d.customWidth:
            out[key] = round(float(d.width), 2)
    return out


def _cond_formats(ws, cap: int = 30) -> list:
    out = []
    try:
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                out.append({"cells": str(cf.sqref), "type": rule.type or "",
                            "operator": rule.operator or "", "formula": list(rule.formula or []),
                            "text": rule.text or ""})
                if len(out) >= cap:
                    return out
    except Exception:  # an unusual rule type openpyxl can't describe
        pass
    return out


def _tables(ws) -> list:
    out = []
    for t in ws.tables.values():
        cols, calc = [], {}
        for tc in t.tableColumns or []:
            cols.append(tc.name)
            f = getattr(tc, "calculatedColumnFormula", None)
            txt = getattr(f, "attr_text", None) or (str(f) if f is not None else "")
            if txt:
                calc[tc.name] = txt if txt.startswith("=") else "=" + txt
        out.append({"name": t.name, "ref": t.ref, "columns": cols, "calculated": calc,
                    "totals_row": bool(t.totalsRowCount)})
    return out


def _print(ws) -> dict:
    out = {}
    try:
        if ws.print_area:
            out["area"] = str(ws.print_area)
        titles = [x for x in (ws.print_title_rows, ws.print_title_cols) if x]
        if titles:
            out["titles"] = ", ".join(str(x) for x in titles)
        if ws.page_setup.orientation:
            out["orientation"] = ws.page_setup.orientation
    except Exception:
        pass
    return out


def _defined_names(wb) -> list:
    out = []
    try:
        items = list(wb.defined_names.items())
    except AttributeError:  # openpyxl < 3.1
        items = [(d.name, d) for d in getattr(wb.defined_names, "definedName", [])]
    for name, d in items:
        out.append({"name": name, "scope": "workbook", "ref": getattr(d, "attr_text", "")})
    for ws in wb.worksheets:
        try:
            for name, d in ws.defined_names.items():
                out.append({"name": name, "scope": ws.title, "ref": getattr(d, "attr_text", "")})
        except AttributeError:
            pass
    return out


def _load(path: str, max_rows: int, max_cols: int):
    """(formula workbook, {sheet: {grid, fmts, comments, unlocked, capped}}) via openpyxl.

    Only cells present in the file are visited (``ws._cells``), so a sheet whose
    formatting runs to column XFD doesn't blow up into millions of empty cells."""
    import openpyxl

    wb_f = openpyxl.load_workbook(path, data_only=False, keep_vba=False)
    wb_v = openpyxl.load_workbook(path, data_only=True, read_only=True, keep_vba=False)
    out = {}
    try:
        for ws in wb_f.worksheets:
            protected = bool(ws.protection.sheet)
            content, comments, unlocked = [], [], []
            for (r, c), cell in ws._cells.items():
                v = cell.value
                if v is not None and v != "":
                    content.append((r, c, v, cell))
                cm = getattr(cell, "comment", None)
                if cm is not None:
                    comments.append((r, c, cm))
                if protected:
                    try:
                        if cell.protection.locked is False:
                            unlocked.append((r - 1, c - 1))
                    except AttributeError:
                        pass
            last_r = max((x[0] for x in content), default=0)
            last_c = max((x[1] for x in content), default=0)
            nr, nc = min(last_r, max_rows), min(last_c, max_cols)
            formulas = [[None] * nc for _ in range(nr)]
            fmts = [[None] * nc for _ in range(nr)]
            for r, c, v, cell in content:
                if r > nr or c > nc:
                    continue
                txt = v if isinstance(v, str) else getattr(v, "text", None)
                if isinstance(txt, str) and txt.startswith("="):
                    formulas[r - 1][c - 1] = txt
                fmts[r - 1][c - 1] = cell.number_format
            values = []
            if nr and nc and ws.title in wb_v.sheetnames:
                for row in wb_v[ws.title].iter_rows(min_row=1, max_row=nr, max_col=nc, values_only=True):
                    row = list(row)[:nc]
                    values.append(row + [None] * (nc - len(row)))
            values += [[None] * nc for _ in range(nr - len(values))]
            out[ws.title] = {
                "grid": G.Grid(ws.title, values, formulas), "fmts": fmts,
                "comments": sorted(comments, key=lambda x: (x[0], x[1])),
                "unlocked": unlocked, "capped": last_r > max_rows or last_c > max_cols,
                "extent": (last_r, last_c),
            }
    finally:
        wb_v.close()
    return wb_f, out


def inspect(path: str, mask: bool = True, max_rows: int = MAX_ROWS) -> dict:
    """The full structure report for one workbook, as a JSON-ready dict."""
    path = os.path.abspath(os.path.expanduser(path))
    _materialize(path)
    wb, loaded = _load(path, max_rows, MAX_COLS)
    buttons = find_buttons(path)
    vba_error = ""
    try:
        modules = V.extract_modules(path) if path.lower().endswith(VBA_EXTS) else []
    except Exception as ex:  # a damaged VBA project, or oletools missing
        modules, vba_error = [], str(ex)
    for m in modules:
        m.code = V.redact_secrets(m.code)

    sheet_of = {}
    for ws in wb.worksheets:
        code = getattr(ws.sheet_properties, "codeName", None)
        if code:
            sheet_of[code] = ws.title
    sheet_of.setdefault(getattr(wb, "code_name", None) or "ThisWorkbook", "(the workbook)")

    sheets, groups = [], {}
    for idx, name in enumerate(wb.sheetnames):
        ws = wb[name]
        base = {"index": idx + 1, "name": name, "state": getattr(ws, "sheet_state", "visible")}
        if name not in loaded:
            sheets.append({**base, "kind": "chartsheet"})
            continue
        d = loaded[name]
        grid = d["grid"]
        tables = _tables(ws)
        blocks = header_blocks(grid, mask=mask, tables=tables)
        patterns, n_formulas = formula_patterns(grid)
        unlocked, n_unlocked = _ranges(d["unlocked"])
        info = {
            **base, "kind": "worksheet",
            "code_name": getattr(ws.sheet_properties, "codeName", None) or "",
            "title": _title(grid, mask),
            "used_range": ws.dimensions, "content_rows": d["extent"][0], "content_cols": d["extent"][1],
            "rows_read": grid.nrows, "capped": d["capped"],
            "freeze_panes": ws.freeze_panes or "",
            "protected": bool(ws.protection.sheet),
            "unlocked_cells": unlocked, "unlocked_more": max(0, n_unlocked - len(unlocked)),
            "tables": tables,
            "merged": [str(m) for m in list(ws.merged_cells.ranges)[:40]],
            "merged_total": len(ws.merged_cells.ranges),
            "validations": [{"type": dv.type or "", "list": dv.formula1 or "", "cells": str(dv.sqref)}
                            for dv in list(ws.data_validations.dataValidation)[:30]],
            "conditional_formats": _cond_formats(ws),
            "print": _print(ws),
            "hidden": _hidden(ws),
            "column_widths": _widths(ws),
            "comments": [{"cell": G.a1(r - 1, c - 1), "author": cm.author or "",
                          "text": show(cm.text, None, mask) if mask else (cm.text or "")}
                         for r, c, cm in d["comments"][:40]],
            "charts": len(getattr(ws, "_charts", []) or []),
            "images": len(getattr(ws, "_images", []) or []),
            "pivots": [getattr(p, "name", "") for p in (getattr(ws, "_pivots", []) or [])],
            "buttons": [b for b in buttons if b["sheet"] == name],
            "formula_cells": n_formulas, "distinct_formulas": len(patterns),
            "header_blocks": [{k: v for k, v in b.items() if k != "_key"} for b in blocks],
        }
        sig = tuple(b["_key"] for b in blocks)
        keys = {p["key"] for p in patterns}
        if sig and sig in groups:
            ref_name, ref_keys = groups[sig]
            info["same_layout_as"] = ref_name
            info["formulas_only_here"] = [{k: v for k, v in p.items() if k != "key"}
                                          for p in patterns if p["key"] not in ref_keys][:60]
            info["formulas_missing_here"] = len(ref_keys - keys)
        else:
            if sig:
                groups[sig] = (name, keys)
            small = grid.nrows <= SMALL_SHEET[0] and grid.ncols <= SMALL_SHEET[1]
            info["cells_scope"] = "all" if small else f"top {TOP_ROWS} rows"
            info["cells"] = cell_rows(grid, mask, grid.nrows if small else TOP_ROWS)
            info["columns"] = column_profiles(grid, d["fmts"], blocks)
            info["formulas"] = [{k: v for k, v in p.items() if k != "key"} for p in patterns[:MAX_PATTERNS]]
            info["formulas_more"] = max(0, len(patterns) - MAX_PATTERNS)
            info["samples"] = [] if small else _samples(grid, blocks, mask)
        sheets.append(info)

    macros = V.catalog(modules, buttons, sheet_of) if modules else {"entry_points": [], "events": [],
                                                                     "helpers": []}
    stat = os.stat(path)
    return {
        "file": os.path.basename(path),
        "file_modified": dt.datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M"),
        "file_size": stat.st_size,
        "generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "tool_version": TOOL_VERSION,
        "values": "masked" if mask else "full",
        "sheet_count": len(wb.sheetnames),
        "sheets": sheets,
        "defined_names": _defined_names(wb),
        "external_links": external_links(path),
        "buttons": buttons,
        "vba": {
            "error": vba_error,
            "modules": [{"name": m.name, "kind": m.kind, "sheet": sheet_of.get(m.name, ""),
                         "description": m.description,
                         "lines": m.code.count("\n") + 1 if m.code else 0, "code": m.code}
                        for m in modules if m.code.strip()],
            **macros,
        },
    }


# ---- markdown -----------------------------------------------------------------

def _md(s) -> str:
    return str(s if s is not None else "").replace("|", "\\|").replace("\n", " ⏎ ")


def _code(s) -> str:
    s = str(s or "").replace("\n", " ")
    if not s:
        return ""
    return f"`` {s} ``" if "`" in s else f"`{s}`"


def _cells_md(rows: list) -> list:
    return [f"- **{r['row']}**: " + " · ".join(f"{c}={_code(v)}" for c, v in r["cells"]) for r in rows]


def to_markdown(rep: dict) -> str:
    out = []
    w = out.append
    vb = rep["vba"]
    w(f"# Workbook structure: {rep['file']}")
    w("")
    w(f"Generated {rep['generated']} by `ledger_app/inspect_workbook.py` (v{rep['tool_version']}) from the "
      f"file saved {rep['file_modified']}. {rep['sheet_count']} sheets, "
      f"{len(vb['entry_points'])} macros you can run, {len(vb['events'])} event macros.")
    w("")
    if rep["values"] == "masked":
        w("Values are **masked**: numbers show as `#`; phone / GSTIN / PAN / account numbers and e-mail "
          "addresses are blurred; password-looking strings in the VBA show as `\"***\"`. Formulas, "
          "headers, labels, number formats and dates are real.")
    else:
        w("Values are **not masked** (`--full-values`).")
    w("")

    w("## Sheets")
    w("")
    w("| # | Sheet | Code name | State | Used range | Rows × cols | Formulas (distinct) | Tables | Buttons | Layout |")
    w("|---|---|---|---|---|---|---|---|---|---|")
    for s in rep["sheets"]:
        if s["kind"] != "worksheet":
            w(f"| {s['index']} | {_md(s['name'])} |  | {s['state']} | chart sheet |  |  |  |  |  |")
            continue
        lay = f"same as {_md(s['same_layout_as'])}" if s.get("same_layout_as") else ""
        w(f"| {s['index']} | {_md(s['name'])} | {_md(s['code_name'])} | {s['state']} | {s['used_range']} | "
          f"{s['content_rows']} × {s['content_cols']} | {s['formula_cells']} ({s['distinct_formulas']}) | "
          f"{_md(', '.join(t['name'] for t in s['tables']))} | {len(s['buttons'])} | {lay} |")
    w("")

    if rep["defined_names"]:
        w("## Defined names")
        w("")
        w("| Name | Scope | Refers to |")
        w("|---|---|---|")
        for d in rep["defined_names"]:
            w(f"| {_md(d['name'])} | {_md(d['scope'])} | {_code(d['ref'])} |")
        w("")
    if rep["external_links"]:
        w("## Links to other workbooks")
        w("")
        for x in rep["external_links"]:
            w(f"- {_code(x)}")
        w("")

    w("## Macros")
    w("")
    if vb["error"]:
        w(f"The VBA could not be read: {vb['error']}")
        w("")
    if vb["entry_points"]:
        w("### Macros you can run (buttons / Alt+F8)")
        w("")
        w("| Macro | Button(s) | MsgBox | InputBox | Other dialogs | What it does |")
        w("|---|---|---|---|---|---|")
        for m in vb["entry_points"]:
            btn = "; ".join(f"“{x['text']}” on {x['sheet']}" for x in m["buttons"])
            w(f"| `{m['name']}` | {_md(btn)} | {m['msgbox']} | {m['inputbox']} | "
              f"{_md(', '.join(m['other_dialogs']))} | {_md(m['description'][:200])} |")
        w("")
        for m in vb["entry_points"]:
            if m["prompts"]:
                w(f"**`{m['name']}` asks:**")
                for p in m["prompts"]:
                    dflt = f" — default {_code(p['default'])}" if p["default"] else ""
                    w(f"- {_md(p['prompt'])}{dflt}")
                w("")
    if vb["events"]:
        w("### Event macros (Excel runs these by itself)")
        w("")
        w("| Macro | Runs on | Sheet | MsgBox | InputBox | What it does |")
        w("|---|---|---|---|---|---|")
        for m in vb["events"]:
            w(f"| `{m['name']}` | {m['event']} | {_md(m['sheet'])} | {m['msgbox']} | {m['inputbox']} | "
              f"{_md(m['description'][:200])} |")
        w("")
    if vb["helpers"]:
        w("### Helper procedures")
        w("")
        for h in vb["helpers"]:
            calls = f" → calls {', '.join(h['calls'])}" if h["calls"] else ""
            desc = f" — {_md(h['description'][:140])}" if h["description"] else ""
            w(f"- `{h['scope']} {h['kind']} {h['name']}({_md(h['params'])})`{desc}{calls}")
        w("")
    if vb["modules"]:
        w("### Modules")
        w("")
        w("| Module | Kind | Sheet | Lines | Notes |")
        w("|---|---|---|---|---|")
        for m in vb["modules"]:
            w(f"| {m['name']} | {m['kind']} | {_md(m['sheet'])} | {m['lines']} | {_md(m['description'][:160])} |")
        w("")
    elif not vb["error"]:
        w("No VBA in this workbook.")
        w("")

    for s in rep["sheets"]:
        if s["kind"] != "worksheet":
            continue
        code = f" (code name `{s['code_name']}`)" if s["code_name"] else ""
        w(f"## Sheet: {s['name']}{code}")
        w("")
        facts = [f"used range `{s['used_range']}`", f"content to row {s['content_rows']}, col "
                 f"{G.col_letter(max(0, s['content_cols'] - 1))}"]
        if s["state"] != "visible":
            facts.append(f"**{s['state']}**")
        if s["freeze_panes"]:
            facts.append(f"frozen at `{s['freeze_panes']}`")
        if s["protected"]:
            facts.append("**protected**")
        if s["capped"]:
            facts.append(f"only the first {s['rows_read']} rows were read")
        w("- " + " · ".join(facts))
        if s["title"]:
            w(f"- Title: {_md(s['title'])}")
        for t in s["tables"]:
            calc = "; ".join(f"{k} {_code(v)}" for k, v in t["calculated"].items())
            w(f"- Excel table **{t['name']}** `{t['ref']}`: {_md(', '.join(t['columns']))}"
              + (f" — calculated: {calc}" if calc else "") + (" — has a totals row" if t["totals_row"] else ""))
        for b in s["buttons"]:
            w(f"- Button “{_md(b['text'])}” runs `{V.button_target(b['macro'])}`")
        for dv in s["validations"]:
            w(f"- Validation ({dv['type']}) on `{dv['cells']}`: {_code(dv['list'])}")
        for cf in s["conditional_formats"]:
            rule = " ".join(x for x in [cf["type"], cf["operator"], cf["text"]] if x)
            fx = ", ".join(cf["formula"])
            w(f"- Conditional format on `{cf['cells']}`: {_md(rule)}" + (f" {_code(fx)}" if fx else ""))
        if s["protected"] and s["unlocked_cells"]:
            more = f" (+{s['unlocked_more']} more)" if s["unlocked_more"] else ""
            w(f"- Unlocked input cells: {', '.join('`%s`' % x for x in s['unlocked_cells'])}{more}")
        if s["merged"]:
            more = f" (+{s['merged_total'] - len(s['merged'])} more)" if s["merged_total"] > len(s["merged"]) else ""
            w(f"- Merged: {', '.join('`%s`' % m for m in s['merged'])}{more}")
        if s["print"]:
            w("- Print: " + " · ".join(f"{k} `{v}`" for k, v in s["print"].items()))
        if s["hidden"]["columns"] or s["hidden"]["rows"]:
            w(f"- Hidden: columns {', '.join(s['hidden']['columns']) or '—'}; rows "
              f"{', '.join(s['hidden']['rows'][:20]) or '—'}")
        extras = [f"{s['charts']} chart(s)" if s["charts"] else "", f"{s['images']} image(s)" if s["images"] else "",
                  f"pivot tables: {', '.join(s['pivots'])}" if s["pivots"] else ""]
        if any(extras):
            w("- " + " · ".join(x for x in extras if x))
        for cm in s["comments"]:
            w(f"- Comment on `{cm['cell']}`: {_md(cm['text'])}")
        w("")
        if s["header_blocks"]:
            w("**Header rows:**")
            w("")
            for hb in s["header_blocks"]:
                rep_ = f" (repeats at rows {', '.join(map(str, hb['repeats'][:10]))})" if hb["repeats"] else ""
                w(f"- Row {hb['row']} `{hb['range']}`{rep_}: " +
                  " · ".join(f"{x['col']}=“{_md(x['label'])}”" for x in hb["labels"]))
            w("")
        if s.get("same_layout_as"):
            w(f"_Same layout as **{s['same_layout_as']}** — see that sheet for the cells, columns and formulas._")
            if s["formulas_only_here"]:
                w("")
                w("Formulas only on this sheet:")
                for f in s["formulas_only_here"]:
                    w(f"- {_code(f['formula'])} — {f['cells']} ({f['count']} cell{'s' if f['count'] != 1 else ''})")
            if s["formulas_missing_here"]:
                w(f"({s['formulas_missing_here']} of that sheet's formulas aren't on this one.)")
            w("")
            continue
        if s.get("cells"):
            w(f"**Cells ({s['cells_scope']}):**")
            w("")
            out += _cells_md(s["cells"])
            w("")
        if s.get("columns"):
            w("**Columns:**")
            w("")
            w("| Col | Header | Filled | Types | Format | First formula | Last formula | Last row |")
            w("|---|---|---|---|---|---|---|---|")
            for c in s["columns"]:
                types = ", ".join(f"{k} {v}" for k, v in c["types"].items())
                w(f"| {c['col']} | {_md(c['header'])} | {c['filled']} | {types} | "
                  f"{_code(_md(c['number_format']))} | {_code(_md(c['first_formula']))} | "
                  f"{_code(_md(c['last_formula']))} | {c['last_row']} |")
            w("")
        if s.get("formulas"):
            w(f"**Formulas** ({s['distinct_formulas']} distinct in {s['formula_cells']} cells):")
            w("")
            for f in s["formulas"]:
                w(f"- {_code(f['formula'])} — {f['cells']}" + (f" ({f['count']} cells)" if f["count"] > 1 else ""))
            if s["formulas_more"]:
                w(f"- … and {s['formulas_more']} more distinct formulas (see the .json)")
            w("")
        if s.get("samples"):
            w("**Sample rows** (under the first header row):")
            w("")
            out += _cells_md(s["samples"])
            w("")

    if vb["modules"]:
        w("## VBA source")
        w("")
        for m in vb["modules"]:
            where = f", sheet “{m['sheet']}”" if m["sheet"] else ""
            w(f"### {m['name']} ({m['kind']} module{where}, {m['lines']} lines)")
            w("")
            w("```vba")
            w(m["code"])
            w("```")
            w("")
    return "\n".join(out).rstrip() + "\n"


# ---- files --------------------------------------------------------------------

def slug(path: str) -> str:
    stem = os.path.splitext(os.path.basename(path))[0]
    return re.sub(r"[^a-z0-9]+", "-", stem.lower()).strip("-") or "workbook"


def write_report(path: str, out_dir: str = OUT_DIR, mask: bool = True, max_rows: int = MAX_ROWS) -> dict:
    rep = inspect(path, mask=mask, max_rows=max_rows)
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.join(out_dir, slug(path))
    with open(base + ".json", "w", encoding="utf-8") as fh:
        json.dump(rep, fh, indent=1, ensure_ascii=False, default=str)
    with open(base + ".md", "w", encoding="utf-8") as fh:
        fh.write(to_markdown(rep))
    return {"md": base + ".md", "json": base + ".json", "sheets": rep["sheet_count"],
            "macros": len(rep["vba"]["entry_points"]), "events": len(rep["vba"]["events"]),
            "modules": len(rep["vba"]["modules"]), "vba_error": rep["vba"]["error"]}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Write a structure report (sheets, formulas, VBA) for workbooks")
    ap.add_argument("workbooks", nargs="*",
                    help="paths to .xlsm / .xlsx files (default: find Master Ledger.xlsm)")
    ap.add_argument("--out", default=OUT_DIR, help="output folder (default: ledger_app/out/)")
    ap.add_argument("--manifest", help="also list the files written, one per line, in this file")
    ap.add_argument("--full-values", action="store_true", help="don't mask numbers and personal details")
    ap.add_argument("--max-rows", type=int, default=MAX_ROWS, help="rows read per sheet")
    args = ap.parse_args(argv)

    paths = [os.path.expanduser(p) for p in args.workbooks]
    if not paths:
        found = find_ledger()
        if not found:
            print(f"Couldn't find {LEDGER_NAME}. Pass its path, e.g.\n"
                  f"  bash ledger_app/report.sh \"/path/to/{LEDGER_NAME}\"", file=sys.stderr)
            return 1
        print(f"Using {found}")
        paths = [found]
    for p in paths:
        if not os.path.exists(p):
            print(f"not found: {p}", file=sys.stderr)
            return 1
        print(f"Reading {os.path.basename(p)} … (a big workbook takes a minute)")
        res = write_report(p, args.out, mask=not args.full_values, max_rows=args.max_rows)
        print(f"  {res['sheets']} sheets · {res['modules']} VBA modules · {res['macros']} macros · "
              f"{res['events']} event macros")
        if res["vba_error"]:
            print(f"  ⚠️  VBA not read: {res['vba_error']}", file=sys.stderr)
        print(f"  -> {os.path.relpath(res['md'])}\n  -> {os.path.relpath(res['json'])}")
        if args.manifest:
            with open(args.manifest, "a", encoding="utf-8") as fh:
                fh.write(f"{res['md']}\n{res['json']}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
