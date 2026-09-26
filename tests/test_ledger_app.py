"""Tests for the Master Ledger structure report (``ledger_app``).

No real ledger is needed: workbooks are built with openpyxl, VBA is parsed
from source strings, and the macro buttons come from a hand-made package.
What's covered is what the report promises — masking, header and formula
detection, layout grouping, and the VBA catalog — plus the safety of
``report.sh --push``: the report reaches only the private repo, this repo is
left exactly as it was, and a public target is refused. All data is made up.
"""

import datetime as dt
import json
import os
import plistlib
import shutil
import subprocess
import sys
import time
import zipfile

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.comments import Comment                       # noqa: E402
from openpyxl.formatting.rule import CellIsRule             # noqa: E402
from openpyxl.styles import PatternFill, Protection         # noqa: E402
from openpyxl.workbook.defined_name import DefinedName      # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.table import Table                  # noqa: E402

from ledger_app import grid as G                  # noqa: E402
from ledger_app import inspect_workbook as I      # noqa: E402
from ledger_app import vba as V                   # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


# ---- grid helpers -------------------------------------------------------------

def test_col_letters_round_trip():
    for c in (0, 25, 26, 51, 701, 702, 16383):
        assert G.col_index(G.col_letter(c)) == c
    assert G.col_letter(26) == "AA" and G.a1(4, 1) == "B5"


def test_grid_is_ragged_safe():
    g = G.Grid("S", [[1], [1, 2, 3], None, []])
    assert (g.nrows, g.ncols) == (4, 3)
    assert g.v(1, 2) == 3 and g.v(2, 0) is None and g.v(9, 9) is None


def test_to_date_reads_dates_not_labels():
    assert G.to_date("26/09/2026") == dt.date(2026, 9, 26)
    assert G.to_date("26.09.26") == dt.date(2026, 9, 26)
    assert G.to_date(dt.datetime(2026, 9, 26, 10)) == dt.date(2026, 9, 26)
    assert G.to_date("Customer") is None and G.to_date("") is None


# ---- VBA ----------------------------------------------------------------------

MODULE1 = '''Attribute VB_Name = "Module1"
Option Explicit

'==========================================
'  Billing helpers
'==========================================

Private Const PROT_PWD As String = "hunter2"   ' sheet protection password, "" if none

'------------------------------------------
' MAIN: save the bill as PDF and log it
'------------------------------------------
Public Sub PrintBill()
    Dim ans As VbMsgBoxResult
    If Range("B13").Value = "" Then MsgBox "Bill To is blank.", vbExclamation: Exit Sub
    ans = MsgBox("Replace the log entry?", vbYesNo + vbDefaultButton2, "Log")
    SaveIt ActiveSheet
    Debug.Print "MsgBox in a string is not a dialog" ' MsgBox in a comment neither
End Sub

Sub Reprint()
    Dim cust As String, sFrom As String, f As Variant
    cust = Trim$(InputBox("Customer name (blank = all)." & vbCrLf & _
                 "Part of the name is enough.", "Reprint"))
    sFrom = InputBox("From date (dd/mm/yyyy):", "Reprint", Format$(Date - 30, "dd/mm/yyyy"))
    f = Application.GetSaveAsFilename()
End Sub

Public Sub NeedsArgs(ws As Worksheet, _
                     Optional n As Long = 1)
End Sub

Private Function SaveIt(ws As Worksheet) As Boolean
    ws.Unprotect Password:="hunter2"
    MsgBox "Saved"
    SaveIt = True
End Function
'''

SHEET3 = '''Attribute VB_Name = "Sheet3"
Attribute VB_Base = "0{00020820-0000-0000-C000-000000000046}"
Attribute VB_PredeclaredId = True
' Recalculate the bill whenever the quantity changes
Private Sub Worksheet_Change(ByVal Target As Range)
    If Target.Address = "$E$20" Then MsgBox "Quantity changed"
End Sub
'''


def _mods():
    return [V.module_from_source("Module1.bas", MODULE1), V.module_from_source("Sheet3.cls", SHEET3)]


def test_module_from_source_strips_attributes_and_detects_kind():
    m1, s3 = _mods()
    assert (m1.name, m1.kind) == ("Module1", "standard")
    assert (s3.name, s3.kind) == ("Sheet3", "document")
    assert "Attribute" not in m1.code and "Attribute" not in s3.code
    assert m1.description == "Billing helpers"


def test_parse_procs_scopes_params_and_descriptions():
    procs = {p.name: p for p in V.parse_procs(_mods()[0])}
    assert list(procs) == ["PrintBill", "Reprint", "NeedsArgs", "SaveIt"]
    assert procs["Reprint"].scope == "Public"                  # no keyword = Public
    assert procs["SaveIt"].scope == "Private" and procs["SaveIt"].kind == "Function"
    assert procs["NeedsArgs"].required_params == 1              # Optional doesn't count
    assert procs["NeedsArgs"].params == "ws As Worksheet, Optional n As Long = 1"
    assert procs["PrintBill"].description == "MAIN: save the bill as PDF and log it"


def test_dialogs_ignore_strings_and_comments():
    procs = {p.name: p for p in V.parse_procs(_mods()[0])}
    assert procs["PrintBill"].msgbox == 2
    rp = procs["Reprint"]
    assert rp.inputbox == 2
    assert rp.prompts[0]["prompt"] == "Customer name (blank = all).\nPart of the name is enough."
    assert rp.prompts[0]["title"] == "Reprint"
    assert rp.prompts[1]["default"] == 'Format$(Date - 30, "dd/mm/yyyy")'
    assert "save dialog (GetSaveAsFilename)" in rp.other_dialogs


def test_catalog_rolls_up_calls_and_maps_buttons_and_events():
    cat = V.catalog(_mods(), [{"sheet": "Bill", "text": "Print", "macro": "[0]!PrintBill"}],
                    {"Sheet3": "Bill"})
    entry = {m["proc"]: m for m in cat["entry_points"]}
    assert list(entry) == ["PrintBill", "Reprint"]             # macro with a button first
    assert entry["PrintBill"]["buttons"] == [{"sheet": "Bill", "text": "Print"}]
    assert entry["PrintBill"]["msgbox"] == 3                    # 2 + the one in SaveIt
    assert entry["PrintBill"]["calls"] == ["Module1.SaveIt"]
    [ev] = cat["events"]
    assert (ev["name"], ev["event"], ev["sheet"]) == ("Sheet3.Worksheet_Change", "Change", "Bill")
    assert {h["name"] for h in cat["helpers"]} == {"Module1.NeedsArgs", "Module1.SaveIt"}


def test_redact_secrets_blanks_passwords_only():
    red = V.redact_secrets(MODULE1)
    assert "hunter2" not in red
    assert 'PROT_PWD As String = "***"   \' sheet protection password, "" if none' in red
    assert 'ws.Unprotect Password:="***"' in red
    assert 'MsgBox "Saved"' in red and 'MsgBox "Bill To is blank."' in red


def test_code_only_blanks_strings_and_drops_comments():
    assert V.code_only('MsgBox "a ""b"" c" \' MsgBox') == 'MsgBox "" '
    assert V.code_only("x = 1 : Rem note") == "x = 1 : "


def test_button_target():
    assert V.button_target("[0]!PrintBillPDFs") == "PrintBillPDFs"
    assert V.button_target("'Master Ledger.xlsm'!Module1.Refresh") == "Module1.Refresh"


# ---- formulas / masking / headers ---------------------------------------------------

def test_to_r1c1_makes_filled_down_formulas_equal():
    assert I.to_r1c1("=B5*2", 4, 2) == "=RC[-1]*2"
    assert I.to_r1c1("=E20*F20", 19, 6) == I.to_r1c1("=E21*F21", 20, 6)
    assert I.to_r1c1("=$B$13+B$1", 4, 1) == "=R13C2+R1C"
    assert I.to_r1c1('=IF(A1="B5","x",A1)', 0, 1) == '=IF(RC[-1]="B5","x",RC[-1])'
    assert I.to_r1c1("='My Sheet'!A1+Sheet2!B2", 0, 0) == "='My Sheet'!RC+Sheet2!R[1]C[1]"
    assert I.to_r1c1("=LOG10(A1)+ATAN2(1,2)+Table1[Qty]", 0, 0) == "=LOG10(RC)+ATAN2(1,2)+Table1[Qty]"


def test_formula_patterns_group_a_filled_down_column():
    g = G.Grid("S", [[None, None]] + [[r, f"=A{r + 1}*2"] for r in range(1, 5)] + [[None, "=SUM(B2:B5)"]],
               [[None, None]] + [[None, f"=A{r + 1}*2"] for r in range(1, 5)] + [[None, "=SUM(B2:B5)"]])
    pats, total = I.formula_patterns(g)
    assert total == 5
    assert pats[0] == {"key": "=RC[-1]*2", "formula": "=A2*2", "count": 4, "cells": "B2:B5"}
    assert pats[1]["cells"] == "B6" and pats[1]["count"] == 1


def test_mask_text_blurs_personal_details():
    assert I.mask_text("Mob: 9876543210") == "Mob: 9999999999"
    assert I.mask_text("GSTIN: 21ABCDE1234F1Z5") == "GSTIN: 99ABCDE9999F9Z9"
    assert I.mask_text("PAN ABCDE1234F") == "PAN ABCDE9999F"
    assert I.mask_text("mail accounts@example.com") == "mail a***@example.com"
    assert I.mask_text("Invoice 12 of 2026") == "Invoice 12 of 2026"


def test_header_blocks_skip_data_rows_but_keep_group_headers():
    d = dt.datetime(2026, 9, 1)
    log = G.Grid("Log", [
        ["VRIDDHI FUELS"],
        [],
        ["Date", "Bill No", "Vehicle", "Customer", "Product"],
        [d, 101, "Truck A", "Alpha Traders", "HSD"],
        [d, 102, "Truck B", "Beta Minerals", "MS"],
    ])
    assert [b["row"] for b in I.header_blocks(log)] == [3]

    stock = G.Grid("Stock", [
        [None, "Stock", None, None, "Sales"],
        ["Date", "Opening", "Receipt", "Closing", "Qty", "Amount"],
        [d, 1, 2, 3, 4, 5],
    ])
    assert [b["row"] for b in I.header_blocks(stock)] == [2]


def test_header_blocks_ignore_table_rows_and_mask_labels():
    g = G.Grid("T", [["Company", "Address", "Bank"], [None], ["ABC Ltd", "Main Road", "HDFC"]])
    assert [b["row"] for b in I.header_blocks(g, tables=[{"ref": "A1:C3"}])] == [1]
    assert [b["row"] for b in I.header_blocks(g)] == [1, 3]      # without the table it's ambiguous
    ids = G.Grid("I", [["GSTIN: 21ABCDE1234F1Z5", "Account 12345678901234", "Name", "City"]])
    assert I.header_blocks(ids) == []                            # IDs aren't headers


def test_header_blocks_stay_fast_on_big_sheets():
    rows = [["Date", "Customer", "Product", "Remarks"]] + [["x", "Customer", "HSD", "ok"]] * 5000
    t = time.perf_counter()
    I.header_blocks(G.Grid("Big", rows))
    assert time.perf_counter() - t < 3


# ---- package parts ---------------------------------------------------------------

def _package(path):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("xl/workbook.xml",
                   '<workbook xmlns:r="r"><sheets><sheet name="Bill &amp; Log" sheetId="1" r:id="rId1"/>'
                   "</sheets></workbook>")
        z.writestr("xl/_rels/workbook.xml.rels",
                   '<Relationships><Relationship Id="rId1" Type="x/worksheet" Target="worksheets/sheet1.xml"/>'
                   "</Relationships>")
        z.writestr("xl/worksheets/sheet1.xml", "<worksheet/>")
        z.writestr("xl/worksheets/_rels/sheet1.xml.rels",
                   '<Relationships><Relationship Id="rId1" Type="x/vmlDrawing" Target="../drawings/v1.vml"/>'
                   '<Relationship Id="rId2" Type="x/drawing" Target="../drawings/drawing1.xml"/></Relationships>')
        z.writestr("xl/drawings/v1.vml",
                   '<xml><v:shape id="b1"><v:textbox><div><font><b>Print Invoice</b></font></div></v:textbox>'
                   '<x:ClientData ObjectType="Button"><x:FmlaMacro>[0]!PrintBillPDFs</x:FmlaMacro>'
                   "</x:ClientData></v:shape></xml>")
        z.writestr("xl/drawings/drawing1.xml",
                   '<xdr:wsDr><xdr:twoCellAnchor><xdr:sp macro="[0]!Module1.Refresh"><xdr:nvSpPr>'
                   '<xdr:cNvPr id="2" name="Rounded Rectangle 1"/></xdr:nvSpPr><xdr:txBody><a:p><a:r>'
                   "<a:t>Refresh</a:t></a:r></a:p></xdr:txBody></xdr:sp></xdr:twoCellAnchor></xdr:wsDr>")
        z.writestr("xl/externalLinks/_rels/externalLink1.xml.rels",
                   '<Relationships><Relationship Id="rId1" Type="x/externalLinkPath" '
                   'Target="file:///Users/x/Master%20Ledger.xlsm" TargetMode="External"/></Relationships>')


def test_find_buttons_reads_form_controls_and_shapes(tmp_path):
    p = tmp_path / "book.xlsm"
    _package(p)
    assert I.find_buttons(str(p)) == [
        {"sheet": "Bill & Log", "macro": "[0]!PrintBillPDFs", "text": "Print Invoice", "kind": "Button"},
        {"sheet": "Bill & Log", "macro": "[0]!Module1.Refresh", "text": "Refresh", "kind": "shape"},
    ]
    assert I.external_links(str(p)) == ["file:///Users/x/Master Ledger.xlsm"]


def test_find_ledger_uses_the_payment_agents_path(tmp_path, monkeypatch):
    ledger = tmp_path / "Master Ledger.xlsm"
    ledger.write_bytes(b"x")
    plist = tmp_path / "agent.plist"
    plist.write_bytes(plistlib.dumps({"EnvironmentVariables": {"MASTER_LEDGER_PATH": str(ledger)}}))
    monkeypatch.delenv("MASTER_LEDGER_PATH", raising=False)
    monkeypatch.setattr(I, "AGENT_PLIST", str(plist))
    assert I.find_ledger() == str(ledger)
    monkeypatch.setenv("MASTER_LEDGER_PATH", str(ledger))
    monkeypatch.setattr(I, "AGENT_PLIST", str(tmp_path / "missing.plist"))
    assert I.find_ledger() == str(ledger)


# ---- the whole report ----------------------------------------------------------------

def _ledger(path):
    wb = openpyxl.Workbook()
    out = wb.active
    out.title = "Outstanding"
    out.append([None, "Customer", "Sheet", "Balance"])
    out.append([None, "Alpha Traders", "ALPHA", "=ALPHA!G6"])
    out.append([None, "Beta Minerals", "BETA", "=BETA!G6"])
    for name in ("ALPHA", "BETA"):
        ws = wb.create_sheet(name)
        ws["A1"] = "VRIDDHI FUELS"
        ws["A2"] = f"Ledger of {name} · Mob 9876543210"
        for c, h in enumerate(["Date", "Volume", "Price", "Amount", "Paid", "Product", "Balance"], 1):
            ws.cell(4, c, h)
        ws["A5"], ws["B5"], ws["C5"], ws["D5"] = dt.datetime(2026, 9, 1), 12000, 90.5, "=B5*C5"
        ws["F5"], ws["G5"] = "HSD", "=D5-E5"
        ws["A6"], ws["B6"], ws["C6"], ws["D6"] = dt.datetime(2026, 9, 2), 8000, 90.5, "=B6*C6"
        ws["E6"], ws["G6"] = 500000, "=G5+D6-E6"
        ws.freeze_panes = "A5"
        if name == "BETA":
            ws["I2"] = "=SUM(D:D)"                     # a formula only this ledger has
    paid = wb.create_sheet("Master Paid")
    paid.append(["Date", "Customer", "Amount Paid", "Payment Mode", "Remarks"])
    paid.append([dt.datetime(2026, 9, 3), "Alpha Traders", 250000, "HDFC NEFT", "Advance"])
    paid.append([dt.datetime(2026, 9, 4), "Beta Minerals", 120000, "HDFC RTGS", "Part"])
    paid.add_table(Table(displayName="MasterPaid", ref="A1:E3"))
    dv = DataValidation(type="list", formula1='"HDFC NEFT,HDFC RTGS"')
    paid.add_data_validation(dv)
    dv.add("D2:D3")
    paid.conditional_formatting.add("C2:C3", CellIsRule(operator="greaterThan", formula=["100000"],
                                                        fill=PatternFill("solid", fgColor="FFC7CE")))
    paid["E2"].comment = Comment("Call 9876543210 before posting", "Accounts")
    paid.column_dimensions["E"].hidden = True
    bill = wb.create_sheet("Bill")
    bill["B2"], bill["B3"], bill["F2"], bill["G2"] = "Bill To", "Alpha Traders", "Date:", "=TODAY()"
    for ref in ("B3", "C5", "D5"):
        bill[ref].protection = Protection(locked=False)
    bill.protection.sheet = True
    bill.print_area = "A1:H20"
    wb.defined_names["HSD_RSP"] = DefinedName("HSD_RSP", attr_text="Bill!$G$2")
    wb.save(path)


@pytest.fixture
def ledger(tmp_path):
    p = tmp_path / "Master Ledger.xlsx"
    _ledger(p)
    return str(p)


def test_report_covers_every_sheet(ledger):
    rep = I.inspect(ledger)
    s = {x["name"]: x for x in rep["sheets"]}
    assert rep["sheet_count"] == 5 and list(s) == ["Outstanding", "ALPHA", "BETA", "Master Paid", "Bill"]

    alpha = s["ALPHA"]
    assert alpha["freeze_panes"] == "A5" and alpha["cells_scope"] == "all"
    assert [b["row"] for b in alpha["header_blocks"]] == [4]
    assert {"formula": "=B5*C5", "count": 2, "cells": "D5:D6"} in alpha["formulas"]
    assert "Mob 9999999999" in alpha["title"]
    assert ["B", "#"] in next(r["cells"] for r in alpha["cells"] if r["row"] == 5)

    beta = s["BETA"]
    assert beta["same_layout_as"] == "ALPHA" and "cells" not in beta
    assert [f["formula"] for f in beta["formulas_only_here"]] == ["=SUM(D:D)"]

    mp = s["Master Paid"]
    assert mp["tables"][0]["name"] == "MasterPaid"
    assert mp["tables"][0]["columns"] == ["Date", "Customer", "Amount Paid", "Payment Mode", "Remarks"]
    assert mp["validations"] == [{"type": "list", "list": '"HDFC NEFT,HDFC RTGS"', "cells": "D2:D3"}]
    assert mp["conditional_formats"][0]["cells"] == "C2:C3"
    assert mp["comments"] == [{"cell": "E2", "author": "Accounts", "text": "Call 9999999999 before posting"}]
    assert mp["hidden"]["columns"] == ["E"]
    assert [b["row"] for b in mp["header_blocks"]] == [1]

    bill = s["Bill"]
    assert bill["protected"] and bill["unlocked_cells"] == ["B3", "C5:D5"]
    assert "$A$1:$H$20" in bill["print"]["area"]

    assert {"name": "HSD_RSP", "scope": "workbook", "ref": "Bill!$G$2"} in rep["defined_names"]
    assert rep["vba"]["modules"] == [] and rep["vba"]["error"] == ""
    assert "9876543210" not in json.dumps(rep, default=str)
    assert "12000" not in json.dumps(s["ALPHA"], default=str)


def test_full_values_keeps_numbers(ledger):
    alpha = next(x for x in I.inspect(ledger, mask=False)["sheets"] if x["name"] == "ALPHA")
    assert ["B", "12000"] in next(r["cells"] for r in alpha["cells"] if r["row"] == 5)


def test_write_report_writes_markdown_and_json(ledger, tmp_path):
    res = I.write_report(ledger, str(tmp_path / "out"))
    assert res["md"].endswith("master-ledger.md") and res["sheets"] == 5
    md = open(res["md"], encoding="utf-8").read()
    for want in ("## Sheet: ALPHA", "_Same layout as **ALPHA**", "`=B5*C5` — D5:D6 (2 cells)",
                 "Unlocked input cells: `B3`, `C5:D5`", "Excel table **MasterPaid**"):
        assert want in md
    assert json.load(open(res["json"], encoding="utf-8"))["file"] == "Master Ledger.xlsx"


def test_report_on_a_real_macro_workbook():
    """The committed Tanker Billing workbook: VBA, buttons and redaction end to end."""
    pytest.importorskip("oletools")
    path = os.path.join(ROOT, "Tanker Billing.xlsm")
    if not os.path.exists(path):
        pytest.skip("Tanker Billing.xlsm not in the checkout")
    rep = I.inspect(path)
    assert rep["vba"]["modules"] and rep["vba"]["entry_points"]
    assert any(m["buttons"] for m in rep["vba"]["entry_points"])
    assert all(m["sheet"] for m in rep["vba"]["modules"] if m["kind"] == "document")


# ---- report.sh: sends only to a private repo, never touches this one -----------

def _git(*args, cwd=None):
    return subprocess.run(["git", *args], cwd=cwd, capture_output=True, text=True, check=True).stdout


@pytest.fixture
def shell_tools():
    pytest.importorskip("oletools")          # so report.sh never pip-installs during a test
    if not (shutil.which("bash") and shutil.which("git")):
        pytest.skip("needs bash and git")
    if not os.path.isdir(os.path.join(ROOT, ".git")):
        pytest.skip("not a git checkout")


def _run_report(tmp_path, ledger, **env):
    e = {**os.environ, "PYTHON": sys.executable, "LEDGER_REPORT_OUT": str(tmp_path / "out")}
    e.pop("LEDGER_REPORT_URL", None)      # only the test decides where a push goes
    e.update(env)
    return subprocess.run(["bash", os.path.join(ROOT, "ledger_app", "report.sh"), ledger, "--push"],
                          cwd=ROOT, env=e, capture_output=True, text=True, timeout=180)


def test_report_sh_push_sends_only_to_the_private_repo(tmp_path, ledger, shell_tools):
    private = tmp_path / "private.git"
    _git("init", "-q", "--bare", str(private))
    status, head = _git("status", "--porcelain", cwd=ROOT), _git("rev-parse", "HEAD", cwd=ROOT)
    branches = _git("branch", "-a", cwd=ROOT)

    r = _run_report(tmp_path, ledger, LEDGER_REPORT_URL=str(private))

    assert r.returncode == 0, r.stdout + r.stderr
    assert "Report sent to the private repo" in r.stdout
    assert _git("--git-dir", str(private), "ls-tree", "-r", "--name-only", "main").split() == \
        ["master-ledger.json", "master-ledger.md"]
    assert sorted(os.listdir(tmp_path / "out")) == ["master-ledger.json", "master-ledger.md"]
    # This repo is exactly as it was: same files, same commit, same branches.
    assert _git("status", "--porcelain", cwd=ROOT) == status
    assert _git("rev-parse", "HEAD", cwd=ROOT) == head
    assert _git("branch", "-a", cwd=ROOT) == branches


@pytest.mark.parametrize("answer, message", [("200", "is PUBLIC"), ("000", "Couldn't confirm")])
def test_report_sh_refuses_unless_the_repo_is_private(tmp_path, ledger, shell_tools, answer, message):
    fake = tmp_path / "bin"
    fake.mkdir()
    (fake / "curl").write_text(f"#!/bin/sh\necho {answer}\n")   # GitHub's answer, no network
    (fake / "curl").chmod(0o755)
    r = _run_report(tmp_path, ledger, LEDGER_REPORT_REPO="someone/some-repo",
                    PATH=f"{fake}{os.pathsep}{os.environ['PATH']}")
    assert r.returncode == 1 and message in r.stdout and "Nothing was sent" in r.stdout


def test_local_reports_are_git_ignored(shell_tools):
    assert subprocess.run(["git", "check-ignore", "-q", "ledger_app/out/master-ledger.md"],
                          cwd=ROOT).returncode == 0
